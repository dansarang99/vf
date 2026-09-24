import pptx
import json
import os

prs = pptx.Presentation(r"C:\Users\note\Downloads\Web_마지영전시회.pptx")
data = []
img_dir = r"C:\Users\note\vf\vf01_webclone\src\assets\images"
os.makedirs(img_dir, exist_ok=True)

img_count = 0
for s_idx, slide in enumerate(prs.slides):
    slide_info = {'slide': s_idx + 1, 'texts': [], 'images': []}
    for sh_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    slide_info['texts'].append(t)
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            img_count += 1
            ext = 'png'
            if 'jpeg' in shape.image.content_type or 'jpg' in shape.image.content_type:
                ext = 'jpg'
            elif 'webp' in shape.image.content_type:
                ext = 'webp'
            fname = f"slide_{s_idx+1}_img_{sh_idx+1}_{img_count}.{ext}"
            fpath = os.path.join(img_dir, fname)
            with open(fpath, 'wb') as f:
                f.write(shape.image.blob)
            slide_info['images'].append({
                'name': shape.name,
                'file': fname,
                'width': shape.width,
                'height': shape.height,
                'bytes': len(shape.image.blob)
            })
    data.append(slide_info)

with open(r"C:\Users\note\vf\vf01_webclone\src\pptx_dump.json", "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print(f"Extracted slides: {len(data)}")
print(f"Extracted images count: {img_count}")
for d in data:
    if d['texts']:
        print(f"Slide {d['slide']}: {d['texts']}")
