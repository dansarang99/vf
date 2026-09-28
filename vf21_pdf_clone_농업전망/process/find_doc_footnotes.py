import docx
import re
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r"C:\Users\note\vf\vf21_pdf_clone_농업전망\result\[036]_농업전망_제1원본_템플릿_골격_레이아웃.docx")

fn_re = re.compile(r'^(?:\d+\)\s+[가-힣]|\*\s*[가-힣])')
print("=== Scanning Footnotes in [036].docx ===")
for idx, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if fn_re.match(t):
        # find what comes after
        next_t = doc.paragraphs[idx+1].text.strip() if idx+1 < len(doc.paragraphs) else "EOF"
        print(f"P{idx:03d}: '{t[:50]}' -> followed by P{idx+1:03d}: '{next_t[:30]}'")
