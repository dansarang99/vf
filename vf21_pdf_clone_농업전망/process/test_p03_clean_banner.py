import docx
from docx.shared import Inches, Pt, Mm
import win32com.client
import pythoncom
import fitz

docx_path = r"C:\Users\note\vf\vf21_pdf_clone_농업전망\result\[036]_농업전망_제1원본_템플릿_골격_레이아웃.docx"
test_docx = r"C:\Users\note\vf\vf21_pdf_clone_농업전망\process\test_p03_clean_banner.docx"
test_pdf = r"C:\Users\note\vf\vf21_pdf_clone_농업전망\process\test_p03_clean_banner.pdf"
banner_img = r"C:\Users\note\vf\vf21_pdf_clone_농업전망\process\sec1_full_header_p03.png"

doc = docx.Document(docx_path)

# P032 is the empty paragraph with 4 floating drawings
# P033 is '1\t쌀'
# P034 is '1.1.\t수급 동향'

# Let's clear drawings from P032
p32 = doc.paragraphs[32]
p32._element.clear_content()

# In P033, replace content with banner_img
p33 = doc.paragraphs[33]
p33._element.clear_content()
run = p33.add_run()
# Original rect width is 400 pt = 141 mm = 5.55 inches
run.add_picture(banner_img, width=Inches(5.55))
p33.paragraph_format.space_before = Pt(45)
p33.paragraph_format.space_after = Pt(12)

# Clear P034 since '1.1. 수급 동향' is in banner_img
p34 = doc.paragraphs[34]
p34._element.clear_content()
p34.paragraph_format.space_before = Pt(0)
p34.paragraph_format.space_after = Pt(0)

# P035 is '1.1.1. 생산'
p35 = doc.paragraphs[35]
p35.paragraph_format.space_before = Pt(6)
p35.paragraph_format.space_after = Pt(6)

doc.save(test_docx)
print("Saved test_p03_clean_banner.docx")

pythoncom.CoInitialize()
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = 0
try:
    wdoc = word.Documents.Open(test_docx)
    wdoc.SaveAs(test_pdf, FileFormat=17)
    wdoc.Close(False)
finally:
    word.Quit()
    pythoncom.CoUninitialize()

print("Exported test_p03_clean_banner.pdf")

pdoc = fitz.open(test_pdf)
print("Page count:", len(pdoc))
pix = pdoc[2].get_pixmap(dpi=150)
out_img = r"C:\Users\note\vf\vf21_pdf_clone_농업전망\process\test_p03_clean_banner.png"
pix.save(out_img)
print("Saved test_p03_clean_banner.png")
