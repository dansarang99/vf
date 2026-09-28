# -*- coding: utf-8 -*-
"""
Prototype for perfect Section 0 replacement and 37-page zero drift validation.
"""
import os
import sys
import docx
from docx.shared import Pt, RGBColor, Mm
from docx.oxml import parse_xml
import win32com.client as win32
import fitz

sys.stdout.reconfigure(encoding='utf-8')

vf21_dir = r"C:\Users\note\vf\vf21_pdf_clone_농업전망"
res_dir = os.path.join(vf21_dir, "result")
proc_dir = os.path.join(vf21_dir, "process")

f014_base = os.path.join(res_dir, "[014]_농업전망_제2장_국내곡물수급동향과전망_20대청사진_완벽복제.docx")
doc = docx.Document(f014_base)

# 1. Unlink Section 1 header so Section 0 header stays only on Page 1
doc.sections[1].header.is_linked_to_previous = False

# 2. Setup Section 0 header with background image
s0 = doc.sections[0]
hp = s0.header.paragraphs[0]
hrun = hp.add_run()
pic = hrun.add_picture(os.path.join(proc_dir, "page1_stitched_background.png"), width=Mm(196.1), height=Mm(266.0))

inline = hrun._element.xpath('.//wp:inline')[0]
extent = inline.xpath('./wp:extent')[0]
docPr = inline.xpath('./wp:docPr')[0]
graphic = inline.xpath('./a:graphic')[0]

anchor_xml = f'''<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
           xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
           distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0"
           behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
    <wp:simplePos x="0" y="0"/>
    <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
    <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
    <wp:extent cx="{extent.get('cx')}" cy="{extent.get('cy')}"/>
    <wp:effectExtent l="0" t="0" r="0" b="0"/>
    <wp:wrapNone/>
    <wp:docPr id="{docPr.get('id')}" name="{docPr.get('name')}"/>
    <wp:cNvGraphicFramePr/>
    {graphic.xml}
</wp:anchor>'''

anchor = parse_xml(anchor_xml)
inline.getparent().replace(inline, anchor)

# 3. Locate Section 0 paragraphs (up to paragraph 13 which contains sectPr)
p13 = doc.paragraphs[13]
sectPr = p13._element.xpath('./w:pPr/w:sectPr')[0]

# Remove paragraphs 0 through 12
for _ in range(13):
    p_elem = doc.paragraphs[0]._element
    p_elem.getparent().remove(p_elem)

# Now doc.paragraphs[0] is the old paragraph 13, which has sectPr!
# We can clear its runs and reuse it as the last paragraph of Section 0,
# and insert the new Page 1 paragraphs before it!
p_sect = doc.paragraphs[0]
for r in p_sect.runs:
    p_sect._element.remove(r._element)

# Helper to insert paragraph before p_sect
def add_p_before(sect_p, space_before=0, space_after=0, left_indent=0, line_spacing=None):
    new_p = doc.add_paragraph()
    sect_p._element.addprevious(new_p._element)
    new_p.paragraph_format.space_before = Pt(space_before)
    new_p.paragraph_format.space_after = Pt(space_after)
    if left_indent:
        new_p.paragraph_format.left_indent = Pt(left_indent)
    if line_spacing:
        new_p.paragraph_format.line_spacing = Pt(line_spacing)
    return new_p

# Add Page 1 Title block
p_ch = add_p_before(p_sect, space_before=28, space_after=0)
r = p_ch.add_run('| 제2장 |')
r.font.name = '맑은 고딕'
r.font.size = Pt(22)
r.font.bold = True

p_tit = add_p_before(p_sect, space_before=18, space_after=0)
r = p_tit.add_run('국내곡물 수급 동향과 전망')
r.font.name = '맑은 고딕'
r.font.size = Pt(25)
r.font.bold = True

p_au = add_p_before(p_sect, space_before=40, space_after=0)
r = p_au.add_run('박한울¹ · 김다정² · 최준혁³ · 김지훈⁴')
r.font.name = '맑은 고딕'
r.font.size = Pt(10.5)
r.font.bold = True

# Add TOC
toc_items = [
    ('1. 쌀', True, 163, 48),
    ('1.1. 수급 동향', False, 175, 2),
    ('1.2. 2026년 전망', False, 175, 2),
    ('1.3. 중장기 전망', False, 175, 2),
    ('2. 콩', True, 163, 14),
    ('2.1. 수급 동향', False, 175, 2),
    ('2.2. 2026년 전망', False, 175, 2),
    ('2.3. 중장기 전망', False, 175, 2),
    ('3. 감자', True, 163, 14),
    ('3.1. 수급 동향', False, 175, 2),
    ('3.2. 2026년 전망', False, 175, 2),
    ('3.3. 중장기 전망', False, 175, 2),
]
for text, bold, left_ind, sp_bef in toc_items:
    p = add_p_before(p_sect, space_before=sp_bef, space_after=0, left_indent=left_ind, line_spacing=12)
    r = p.add_run(text)
    r.font.name = '맑은 고딕'
    r.font.size = Pt(10 if bold else 9)
    r.font.bold = bold

# Footnotes (place inside p_sect or preceding paragraphs)
fn_items = [
    '1 | 한국농촌경제연구원 전문연구원, phu87@krei.re.kr',
    '2 | 한국농촌경제연구원 전문연구원, swetmug@krei.re.kr',
    '3 | 한국농촌경제연구원 연구원, wnsgur3385@krei.re.kr',
    '4 | 한국농촌경제연구원 연구원, jhkim4209@krei.re.kr',
]
for i, text in enumerate(fn_items):
    p = add_p_before(p_sect, space_before=(70 if i == 0 else 1), space_after=0, left_indent=163, line_spacing=10)
    r = p.add_run(text)
    r.font.name = '맑은 고딕'
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

out_test = os.path.join(proc_dir, "test_full_p1_replaced.docx")
doc.save(out_test)
print(f"Saved full docx with replaced Page 1: {out_test}")

word = win32.Dispatch('Word.Application')
word.Visible = False
word.DisplayAlerts = 0
out_pdf = os.path.join(proc_dir, "test_full_p1_replaced.pdf")
try:
    wdoc = word.Documents.Open(os.path.abspath(out_test), ReadOnly=True)
    if os.path.exists(out_pdf):
        os.remove(out_pdf)
    wdoc.ExportAsFixedFormat(os.path.abspath(out_pdf), ExportFormat=17)
    wdoc.Close(False)
    print(f"Exported full test PDF: {out_pdf}")
finally:
    word.Quit()

pdoc = fitz.open(out_pdf)
print(f"Total pages of full document: {len(pdoc)} (Target: 37)")
# Render Page 1 and Page 2
pix1 = pdoc[0].get_pixmap(dpi=200)
pix1.save(os.path.join(proc_dir, "full_page1_200dpi.png"))
pix2 = pdoc[1].get_pixmap(dpi=200)
pix2.save(os.path.join(proc_dir, "full_page2_200dpi.png"))
pdoc.close()
print("Saved full_page1_200dpi.png and full_page2_200dpi.png!")
