# -*- coding: utf-8 -*-
import docx

doc = docx.Document(r"C:\Users\note\vf\vf21_pdf_clone_농업전망\result\[014]_농업전망_제2장_국내곡물수급동향과전망_20대청사진_완벽복제.docx")
print("Total Paragraphs in DOCX:", len(doc.paragraphs))
print("Total Tables in DOCX:", len(doc.tables))

print("\n--- Inspecting first 25 paragraphs ---")
for i, p in enumerate(doc.paragraphs[:25]):
    text = p.text.strip()
    drawings = len(p._element.xpath('.//*[local-name()="drawing"]'))
    blips = len(p._element.xpath('.//*[local-name()="blip"]'))
    pagebreaks = len(p._element.xpath('.//*[local-name()="br" and @*[local-name()="type"]="page"]'))
    print(f"P{i:02d}: text='{text[:35]}' | drawings={drawings} | blips={blips} | pb={pagebreaks}")
