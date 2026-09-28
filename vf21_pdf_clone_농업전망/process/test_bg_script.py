import docx
from docx.shared import Pt, Inches, Mm, RGBColor
from docx.oxml import parse_xml
import win32com.client as win32
import os
import fitz

doc = docx.Document()
section = doc.sections[0]
section.page_width = Mm(196.1)
section.page_height = Mm(266.0)
section.top_margin = Mm(30)
section.bottom_margin = Mm(30)
section.left_margin = Mm(31)
section.right_margin = Mm(31)

section.different_first_page_header_footer = True
header = section.first_page_header
hp = header.paragraphs[0]
hrun = hp.add_run()
pic = hrun.add_picture(r'C:\Users\note\vf\vf21_pdf_clone_농업전망\process\page1_stitched_background.png', width=Mm(196.1), height=Mm(266.0))

inline = hrun._element.xpath('.//wp:inline')[0]
extent = inline.xpath('./wp:extent')[0]
docPr = inline.xpath('./wp:docPr')[0]
graphic = inline.xpath('./a:graphic')[0]

anchor_xml = f'''<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
           xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
           distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0"
           behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
    <wp:simplePos x="0" y="0"/>
    <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
    <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
    <wp:extent cx="{extent.get('cx')}" cy="{extent.get('cy')}"/>
    <wp:effectExtent l="0" t="0" r="0" b="0"/>
    <wp:wrapNone/>
    <wp:docPr id="{docPr.get('id')}" name="{docPr.get('name')}"/>
    <wp:cNvGraphicFramePr/>
    {graphic.xml}
</wp:anchor>'''

anchor = parse_xml(anchor_xml)
inline.getparent().replace(inline, anchor)

# Add Page 1 Title text
p1 = doc.add_paragraph()
p1.paragraph_format.space_before = Pt(28)
p1.paragraph_format.space_after = Pt(0)
r1 = p1.add_run('| 제2장 |')
r1.font.name = '맑은 고딕'
r1.font.size = Pt(22)
r1.font.bold = True

p2 = doc.add_paragraph()
p2.paragraph_format.space_before = Pt(18)
p2.paragraph_format.space_after = Pt(0)
r2 = p2.add_run('국내곡물 수급 동향과 전망')
r2.font.name = '맑은 고딕'
r2.font.size = Pt(25)
r2.font.bold = True

p3 = doc.add_paragraph()
p3.paragraph_format.space_before = Pt(40)
p3.paragraph_format.space_after = Pt(0)
r3 = p3.add_run('박한울¹ · 김다정² · 최준혁³ · 김지훈⁴')
r3.font.name = '맑은 고딕'
r3.font.size = Pt(10.5)
r3.font.bold = True

# Add TOC indented (starts at left_indent = Pt(163), matching 251pt from page left since left_margin is 88pt)
# 251.1 - 87.9 = 163.2 pt
toc_items = [
    ('1. 쌀', True, Pt(163), Pt(48)),
    ('1.1. 수급 동향', False, Pt(175), Pt(2)),
    ('1.2. 2026년 전망', False, Pt(175), Pt(2)),
    ('1.3. 중장기 전망', False, Pt(175), Pt(2)),
    ('2. 콩', True, Pt(163), Pt(14)),
    ('2.1. 수급 동향', False, Pt(175), Pt(2)),
    ('2.2. 2026년 전망', False, Pt(175), Pt(2)),
    ('2.3. 중장기 전망', False, Pt(175), Pt(2)),
    ('3. 감자', True, Pt(163), Pt(14)),
    ('3.1. 수급 동향', False, Pt(175), Pt(2)),
    ('3.2. 2026년 전망', False, Pt(175), Pt(2)),
    ('3.3. 중장기 전망', False, Pt(175), Pt(2)),
]

for text, bold, left_ind, sp_bef in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = left_ind
    p.paragraph_format.space_before = sp_bef
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(12)
    r = p.add_run(text)
    r.font.name = '맑은 고딕'
    r.font.size = Pt(10 if bold else 9)
    r.font.bold = bold

# Footnotes
fn_items = [
    '1 | 한국농촌경제연구원 전문연구원, phu87@krei.re.kr',
    '2 | 한국농촌경제연구원 전문연구원, swetmug@krei.re.kr',
    '3 | 한국농촌경제연구원 연구원, wnsgur3385@krei.re.kr',
    '4 | 한국농촌경제연구원 연구원, jhkim4209@krei.re.kr',
]
for i, text in enumerate(fn_items):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(163)
    p.paragraph_format.space_before = Pt(70 if i == 0 else 1)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(10)
    r = p.add_run(text)
    r.font.name = '맑은 고딕'
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

out_docx = r'C:\Users\note\vf\vf21_pdf_clone_농업전망\process\test_bg.docx'
doc.save(out_docx)
print('Saved test_bg.docx')

word = win32.Dispatch('Word.Application')
word.Visible = False
word.DisplayAlerts = 0
out_pdf = r'C:\Users\note\vf\vf21_pdf_clone_농업전망\process\test_bg.pdf'
try:
    wdoc = word.Documents.Open(os.path.abspath(out_docx), ReadOnly=True)
    if os.path.exists(out_pdf):
        os.remove(out_pdf)
    wdoc.ExportAsFixedFormat(os.path.abspath(out_pdf), ExportFormat=17)
    wdoc.Close(False)
    print('Exported test_bg.pdf')
finally:
    word.Quit()

pdoc = fitz.open(out_pdf)
print('Total pages:', len(pdoc))
pix = pdoc[0].get_pixmap(dpi=200)
out_png = r'C:\Users\note\vf\vf21_pdf_clone_농업전망\process\test_bg_rendered.png'
pix.save(out_png)
pdoc.close()
print('Saved rendered PNG:', out_png)
