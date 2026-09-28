import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r"C:\Users\note\vf\vf21_pdf_clone_농업전망\result\[036]_농업전망_제1원본_템플릿_골격_레이아웃.docx")

print("=== Searching for Section 2 (콩) and Section 3 (감자) headers ===")
for idx, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if '2\t콩' in t or '2 콩' in t or (t.startswith('2') and '콩' in t and len(t) < 15):
        print(f"P{idx:03d} (Sec 2): '{t}'")
        for j in range(max(0, idx-3), min(len(doc.paragraphs), idx+5)):
            print(f"   neighbor P{j:03d}: text='{doc.paragraphs[j].text[:30]}', drawings={len(doc.paragraphs[j]._element.xpath('.//w:drawing'))}")
    if '3\t감자' in t or '3 감자' in t or (t.startswith('3') and '감자' in t and len(t) < 15):
        print(f"P{idx:03d} (Sec 3): '{t}'")
        for j in range(max(0, idx-3), min(len(doc.paragraphs), idx+5)):
            print(f"   neighbor P{j:03d}: text='{doc.paragraphs[j].text[:30]}', drawings={len(doc.paragraphs[j]._element.xpath('.//w:drawing'))}")
