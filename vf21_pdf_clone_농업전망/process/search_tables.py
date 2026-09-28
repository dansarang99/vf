import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r"C:\Users\note\vf\vf21_pdf_clone_농업전망\result\[036]_농업전망_제1원본_템플릿_골격_레이아웃.docx")
for t_idx, table in enumerate(doc.tables):
    txt = table.cell(0, 0).text.strip()
    full_table_text = " ".join(c.text.strip() for row in table.rows for c in row.cells)
    if '감자' in full_table_text or '수급 동향' in full_table_text:
        print(f"Table {t_idx} (rows={len(table.rows)}, cols={len(table.columns)}): text snippet='{full_table_text[:60]}'")
