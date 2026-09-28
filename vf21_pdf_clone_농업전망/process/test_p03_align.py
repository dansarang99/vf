import docx
from docx.shared import Pt, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import win32com.client
import pythoncom
import fitz

docx_path = r"C:\Users\note\vf\vf21_pdf_clone_농업전망\result\[036]_농업전망_제1원본_템플릿_골격_레이아웃.docx"
test_docx = r"C:\Users\note\vf\vf21_pdf_clone_농업전망\process\test_p03_align.docx"
test_pdf = r"C:\Users\note\vf\vf21_pdf_clone_농업전망\process\test_p03_align.pdf"

doc = docx.Document(docx_path)

# In original PDF:
# '1 쌀' is at y = 117.5 pt, x = 117.0 pt.
# Section 2 is Page 3.
# Let's set P033 space_before to 105 pt so it drops down right into the circle!
# And let's check P034 '1.1. 수급 동향': in original PDF it is at y = 218.4 pt.
# P033 height is ~25 pt, so space_after on P033 or space_before on P034 should bridge the gap (218.4 - 142 = ~76 pt).
# And '1.1.1. 생산' is at y = 271.7 pt (271.7 - 235 = ~36 pt).
# And first bullet is at y = 313.8 pt (313.8 - 287 = ~26 pt).

p33 = doc.paragraphs[33]
p33.paragraph_format.space_before = Pt(105)
p33.paragraph_format.space_after = Pt(72)

p34 = doc.paragraphs[34]
p34.paragraph_format.space_before = Pt(0)
p34.paragraph_format.space_after = Pt(32)

p35 = doc.paragraphs[35]
p35.paragraph_format.space_before = Pt(0)
p35.paragraph_format.space_after = Pt(24)

doc.save(test_docx)
print("Saved test_p03_align.docx")

pythoncom.CoInitialize()
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = 0
try:
    wdoc = word.Documents.Open(test_docx)
    wdoc.SaveAs(test_pdf, FileFormat=17)
    wdoc.Close(False)
finally:
    word.Quit()
    pythoncom.CoUninitialize()

print("Exported test_p03_align.pdf")

pdoc = fitz.open(test_pdf)
print("Page count:", len(pdoc))
pix = pdoc[2].get_pixmap(dpi=150)
out_img = r"C:\Users\note\vf\vf21_pdf_clone_농업전망\process\test_p03_aligned.png"
pix.save(out_img)
print("Saved test_p03_aligned.png")
